"""General local document import with format-aware text extraction."""

from __future__ import annotations

import contextlib
import csv
import html as html_lib
import io
import json
import logging
import re
import subprocess
from collections.abc import Iterator
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, ClassVar

from varyz.connectors.base import SkipItem, decode_text, register, resolve_paths
from varyz.domain import Cursor, Kind, NormalizedEntry, RawItem, TimePrecision

log = logging.getLogger(__name__)


_HTML_DROP_RE = re.compile(r"(?is)<(script|style|head)\b.*?</\1>")
_HTML_BREAK_RE = re.compile(r"(?i)<(?:br|/p|/div|/tr|/li|/h[1-6])[^>]*>")
_HTML_TAG_RE = re.compile(r"<[^>]+>")


def strip_html(html: str) -> str:
    text = _HTML_DROP_RE.sub(" ", html)
    text = _HTML_BREAK_RE.sub("\n", text)
    text = _HTML_TAG_RE.sub(" ", text)
    text = html_lib.unescape(text).replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*", "\n\n", text)
    return text.strip()


EXTENSIONS = {
    ".csv",
    ".docx",
    ".html",
    ".htm",
    ".json",
    ".md",
    ".pdf",
    ".pptx",
    ".rtf",
    ".txt",
    ".xlsx",
}
_SKIP_NAMES = {"desktop.ini", "thumbs.db"}
_GIT_STATUS_TIMEOUT = 60

# Directories that hold machine-generated or vendored files. A developer's
# folder is mostly dependencies by volume; walking into them buries real
# documents in lockfiles and third-party READMEs. Pruned during the walk, so
# the cost of a large `node_modules` is not even paid in stat calls.
_SKIP_DIRS = {
    "__pycache__",
    "build",
    "deriveddata",
    "dist",
    "env",
    "node_modules",
    "obj",
    "out",
    "pods",
    "runtime",
    "site-packages",
    "target",
    "vendor",
    "venv",
}

# Generated files that pass the extension filter but carry no meaning.
_SKIP_FILENAMES = {
    "package-lock.json",
    "yarn.lock",
    "pnpm-lock.yaml",
    "composer.lock",
    "poetry.lock",
    "cargo.lock",
    # Machine-learning artefacts. A tokenizer vocabulary is 2,500 chunks of
    # single tokens: it matches almost any query weakly and answers none.
    "vocab.json",
    "tokenizer.json",
    "tokenizer_config.json",
    "special_tokens_map.json",
    "added_tokens.json",
    "generation_config.json",
    "preprocessor_config.json",
    "trainer_state.json",
    "merges.txt",
}

# A directory holding model weights is a model, not a folder of documents —
# whatever JSON and text sits beside the weights describes the tensors.
# Detected by content rather than by name so it holds for any framework.
_WEIGHT_SUFFIXES = {".safetensors", ".gguf", ".bin", ".ckpt", ".onnx", ".pt", ".pth"}
_CHECKPOINT_DIR = re.compile(r"^checkpoint[-_]?\d+$", re.IGNORECASE)

# Beyond this, a file is a dataset being stored rather than a document being
# read. Truncating keeps a long report usable while stopping one 6 MB file
# from producing thousands of chunks.
MAX_DOCUMENT_CHARS = 120_000
_TRUNCATION_MARKER = "\n\n… truncated by varyz …"

# Formats that are equally often a written document and a data dump. Only
# these are tested for tabularity; a .md or .docx is taken at face value.
_AMBIGUOUS_EXTENSIONS = {".csv", ".json", ".txt"}
_TABULAR_PROBE_BYTES = 64_000

# Formats whose bytes are meant to decode straight to text. The rest of
# EXTENSIONS are binary containers that a library unpacks, so a binary-looking
# payload is exactly what they should have.
_PLAIN_TEXT_EXTENSIONS = {".csv", ".html", ".htm", ".json", ".md", ".rtf", ".txt"}


# Fail-closed deny list. These hold personal message archives, credentials, or
# another local tool's private store. They are refused whatever the
# configuration says and whether or not git ignores them — a walk that reaches
# one has already gone wrong. Matched against a directory's own name.
_DENY_DIR_EXACT = {
    "whatsapp_bridge",
    "vault",
    ".secrets",
    "_preview",
    "credentials",
    "secrets",
    # Browser profiles carry cookies, session tokens, and history. varyz does
    # not read browsing history, and the documents connector must never read
    # them as files.
    "chrome-profile",
    "chrome_profile",
    "browser-profile",
    "user data",
    "playwright",
    "puppeteer",
}
_DENY_DIR_PREFIX = ("whatsapp_out", "whatsapp_export")


def _is_denied(directory: Path) -> bool:
    name = directory.name.lower()
    return name in _DENY_DIR_EXACT or name.startswith(_DENY_DIR_PREFIX)


def _holds_model_weights(children: list[Path]) -> bool:
    # Files only: a Windows drive copied over leaves a `$RECYCLE.BIN` folder,
    # and that name alone used to hide a whole Documents folder.
    return any(
        child.suffix.lower() in _WEIGHT_SUFFIXES and child.is_file() for child in children
    )


_NUMERIC_ROW = re.compile(r"^[-+]?[\d.,:eE/\s%$-]*$")

# How wide a record has to be before its regularity is evidence of a machine.
# Measured over this corpus, the ambiguous files that hold a constant number of
# fields split cleanly either side of this: the tables have 6, 7, 14 and 80
# columns, and everything else has one or two — a requirements.txt, a list of
# prompt templates. Set at the wide end of that gap on purpose, because a
# narrow export with a name and an address in it is somebody's address book.
_TABULAR_MIN_FIELDS = 6
_TABULAR_ROW_AGREEMENT = 0.9


def looks_tabular(text: str) -> bool:
    """Is this a table rather than something written?

    Judged on rows, not characters: a prose test cannot separate a tokenizer
    vocabulary (which is all words) from a paragraph, but a row test separates
    a price series from a report without ambiguity. Only consulted for formats
    that can plausibly be either.

    Two shapes, because numbers alone missed the ones that hurt. A GPU
    telemetry log starts every row with the application, the card and the CPU
    before its eighty numeric columns, so a numeric-row test reads it as prose
    — and it retrieves beautifully for "what GPU is this machine running",
    where it outranked the file that answers the question. What gives it away
    is not its content but its regularity: one header and four hundred records
    cut to exactly the same width, which nothing written by a person does.
    """
    rows = [line for line in text.splitlines()[:400] if line.strip()]
    body = rows[1:]
    if len(body) < 10:
        return False
    numeric = sum(1 for row in body if _NUMERIC_ROW.match(row.strip()))
    if numeric / len(body) > 0.8:
        return True
    fields = rows[0].count(",")
    if fields + 1 < _TABULAR_MIN_FIELDS:
        return False
    # Compared against the header rather than to each other, so a file that
    # merely contains a table somewhere is not condemned by it.
    agreeing = sum(1 for row in body if row.count(",") == fields)
    return agreeing / len(body) >= _TABULAR_ROW_AGREEMENT


_BINARY_PROBE_CHARS = 4_000
_MAX_CONTROL_SHARE = 0.10


def looks_binary(text: str) -> bool:
    """Is this decoded text, or a binary file that happens to end in `.md`?

    Found in the index: five `._README.md` files, each a macOS AppleDouble
    resource fork, contributing 21 embedded chunks of structure bytes that
    could be returned as evidence. The name check below catches those by
    itself; this catches the general case, because any binary saved under a
    text extension decodes to something and every decoding produces *a* string.

    Judged after decoding rather than on the raw bytes, which is the whole
    trick: UTF-16 text is half NUL bytes on disk and contains none once
    decoded, so a raw-byte test would refuse exactly the files
    `decode_text` exists to rescue. A NUL surviving the decode is decisive —
    text does not contain them — and a high share of other control characters
    catches the rest without touching documents that merely use tabs.
    """
    probe = text[:_BINARY_PROBE_CHARS]
    if not probe:
        return False
    if "\x00" in probe:
        return True
    control = sum(1 for c in probe if ord(c) < 32 and c not in "\t\r\n")
    return control / len(probe) > _MAX_CONTROL_SHARE


def _is_generated(directory: Path) -> bool:
    """Caches are regenerable by definition and are never source material.

    Not every project gitignores them — persona keeps 5,472 hash-named model
    evaluation caches under a tracked path — so name them explicitly too.
    """
    name = directory.name.lower()
    return name == "cache" or name.endswith(("_cache", ".cache", "-cache"))


def _git_ignored(repo: Path) -> set[Path]:
    """Paths git ignores in `repo`, collapsed at directory level.

    A developer's own .gitignore is the most accurate available statement of
    what is generated, vendored, or private. Honouring it keeps build output,
    virtualenvs, credentials, and other tools' local databases out of the
    index without asking the user to enumerate them.
    """
    try:
        result = subprocess.run(
            ["git", "-C", str(repo), "status", "--ignored", "--porcelain=v1"],
            capture_output=True,
            timeout=_GIT_STATUS_TIMEOUT,
            check=False,
            creationflags=int(getattr(subprocess, "CREATE_NO_WINDOW", 0)),
        )
    except (OSError, subprocess.SubprocessError):
        return set()
    if result.returncode != 0:
        return set()
    ignored: set[Path] = set()
    for line in result.stdout.decode("utf-8", "replace").splitlines():
        if line.startswith("!! "):
            entry = line[3:].strip().strip('"').rstrip("/")
            try:
                ignored.add((repo / entry).resolve())
            except OSError:
                continue
    return ignored


def _walk_documents(root: Path, extra_denied: frozenset[Path] = frozenset()) -> Iterator[Path]:
    """Files under `root`, pruning dot-directories, dependency trees,
    git-ignored paths, and the fail-closed deny list."""
    stack: list[tuple[Path, frozenset[Path]]] = [(root, frozenset())]
    # Directories already walked, by real path. A junction or symlink pointing
    # at an ancestor turns the tree into a cycle, and `is_dir()` follows both:
    # measured on a folder with one `mklink /J` back to its own root, the walk
    # re-entered 64 times and only stopped when Windows refused the path
    # length. These are not exotic — Windows scatters compatibility junctions
    # through a user profile, and a repository symlinked into a workspace is
    # ordinary. With long paths enabled there is nothing to stop it at all.
    seen: set[Path] = set()
    with contextlib.suppress(OSError):
        seen.add(root.resolve())
    while stack:
        directory, ignored = stack.pop()
        if (directory / ".git").exists():
            # Entering a repository: adopt its ignore rules for the subtree.
            ignored = frozenset(_git_ignored(directory))
        try:
            children = sorted(directory.iterdir())
        except OSError:
            continue
        if _holds_model_weights(children):
            continue
        for child in children:
            try:
                if child.is_dir():
                    name = child.name.lower()
                    if child.name.startswith(".") or name in _SKIP_DIRS:
                        continue
                    if _is_denied(child) or _is_generated(child):
                        continue
                    if _CHECKPOINT_DIR.match(child.name):
                        continue
                    resolved = child.resolve()
                    if resolved in ignored or resolved in extra_denied:
                        continue
                    if resolved in seen:
                        continue
                    seen.add(resolved)
                    stack.append((child, ignored))
                elif child.is_file():
                    if child.resolve() in ignored:
                        continue
                    yield child
            except OSError:
                continue


def _may_be_document(path: Path) -> bool:
    """Whether a file's name alone lets it be a document worth reading."""
    name = path.name.lower()
    return not (
        path.suffix.lower() not in EXTENSIONS
        or name in _SKIP_NAMES
        or name in _SKIP_FILENAMES
        or path.name.startswith("~$")
        # AppleDouble sidecars. macOS writes one beside every real file on a
        # non-native volume, so they arrive by the hundred on any shared drive
        # or USB stick, and each is a resource fork wearing the real file's
        # extension.
        or path.name.startswith("._")
    )


@register
class DocumentsConnector:
    connector_id: ClassVar[str] = "documents"
    # Every run walks the configured roots in full and yields every file it
    # still finds, so a path missing from a completed walk is a file that is
    # no longer there — deleted, renamed, moved out, or newly ignored.
    enumerates_corpus: ClassVar[bool] = True

    def __init__(self) -> None:
        # How far through its files this discovery is, for the sync's bar.
        self.read_fraction: float | None = None

    def discover(
        self, options: dict[str, Any], cursor: Cursor | None
    ) -> Iterator[tuple[RawItem, Cursor | None]]:
        if not options.get("approved", False):
            raise ValueError("document folders require explicit approval in Varyz Imports")
        roots = resolve_paths(options, self.connector_id)
        max_bytes = int(float(options.get("max_size_mb", 50)) * 1024 * 1024)
        # Varyz must never index its own store: the vault would grow a copy of
        # itself on every sync.
        extra_denied: set[Path] = set()
        own_data = options.get("_data_dir")
        if own_data:
            with contextlib.suppress(OSError):
                extra_denied.add(Path(str(own_data)).resolve())
        denied = frozenset(extra_denied)
        # Walked in full before any file is read, so the sync can say how far
        # through them it is.
        paths = [
            path
            for root in roots
            for path in ([root] if root.is_file() else _walk_documents(root, denied))
            if _may_be_document(path)
        ]
        for index, path in enumerate(paths):
            self.read_fraction = index / len(paths)
            if not path.is_file():
                continue
            try:
                # One stat, kept. The mtime used to be read again while
                # building the metadata below, outside this guard, so a
                # file that vanished in between — a build artefact, a
                # download, an editor's swap file — raised out of the
                # generator and took the whole sync with it. Indexing a
                # folder somebody is working in is the normal case.
                info = path.stat()
                if info.st_size > max_bytes:
                    continue
                raw = path.read_bytes()
                resolved = str(path.resolve())
            except OSError:
                continue
            suffix = path.suffix.lower()
            if suffix in _PLAIN_TEXT_EXTENSIONS:
                # Decoded whole and then sliced, not sliced and then
                # decoded. A byte prefix can end mid-character, and for a
                # UTF-16 file the fallback then reads the NUL half of every
                # pair as text — which `looks_binary` would correctly call
                # binary, refusing a real document over an artefact of
                # where the probe was cut. Both tests read a few thousand
                # characters; only the decode was ever bounded.
                probe = decode_text(raw)[:_TABULAR_PROBE_BYTES]
                if looks_binary(probe):
                    continue
                if suffix in _AMBIGUOUS_EXTENSIONS and looks_tabular(probe):
                    continue
            yield (
                RawItem(
                    external_id=resolved,
                    payload=raw,
                    payload_ext=path.suffix.lower(),
                    metadata={"path": resolved, "mtime": info.st_mtime},
                ),
                None,
            )

    def still_present(self, external_id: str) -> bool:
        """Was this file merely unreadable this time, rather than gone?

        Consulted before an entry is pruned, and answering only that one
        question. The walk skips a file for two quite different sorts of
        reason, and this separates them: it could not read the file, which is
        usually temporary and usually means another program has it open — or it
        read the file and decided against it, which is a decision that should
        stand.

        Checking existence alone was the wrong test. It kept every entry whose
        file was still on disk, which is exactly the content a user has just
        finished excluding: a folder added to the deny list, a path newly
        git-ignored, a binary that stopped qualifying as a document. "Stop
        indexing this" has to actually remove it, or the setting is decorative
        and the privacy direction is the wrong one.

        So: gone is gone, unreadable is kept, and readable-but-not-offered is
        the walk's decision to make. The remaining error is a file that was
        briefly locked during the walk and free a moment later, which costs one
        re-embed on the next sync and corrects itself.
        """
        path = Path(external_id)
        try:
            if not path.exists():
                return False
            with path.open("rb") as handle:
                handle.read(1)
        except OSError:
            return True
        return False

    def normalize(self, item: RawItem) -> NormalizedEntry:
        path = Path(item.metadata["path"])
        try:
            body = extract_document_text(item.payload_bytes(), item.payload_ext)
        except UnreadableDocument as exc:
            log.warning("skipping %s: %s", path, exc)
            raise SkipItem(item.external_id) from None
        if not body.strip():
            # A file with no text is not a document. Without this the entry was
            # made anyway and chunked to its own filename, so the index carried
            # a searchable `index`, `stdout`, `Guidelines` pointing at nothing.
            raise SkipItem(item.external_id)
        occurred = datetime.fromtimestamp(float(item.metadata["mtime"]), tz=UTC)
        return NormalizedEntry(
            kind=Kind.DOCUMENT,
            external_id=item.external_id,
            title=path.stem,
            body=body.strip(),
            occurred_start=occurred,
            captured_at=occurred,
            time_precision=TimePrecision.DAY,
            significance=0.55,
        )


class UnreadableDocument(Exception):
    """A file whose bytes are not the format its name claims.

    A truncated download, an HTML error page saved as .pdf, a rename somebody
    did to make a file open in the right app. The libraries that unpack the
    binary containers raise on those, and ingest counted the exception as an
    unexpected failure: a twenty-line third-party traceback in the terminal on
    every sync, `connect status` reading errors for as long as the file stayed
    in the folder, and the cursor held back because the run had failed.

    None of that is true of it. The file is unreadable, which is a property of
    the file and will be just as true next time -- so it belongs with the other
    things extraction can only find out by trying, with a line naming it,
    because unlike an empty file it is one its owner may want to know about.
    """


# The containers a library unpacks, which is EXTENSIONS minus the ones whose
# bytes decode straight to text. Only these are wrapped: a failure inside one
# is a statement about the bytes it was handed, where a failure in the text
# path would be ours and should not be turned into a skipped file.
_PARSED_EXTENSIONS = EXTENSIONS - _PLAIN_TEXT_EXTENSIONS


def extract_document_text(raw: bytes, ext: str) -> str:
    if ext in _PARSED_EXTENSIONS:
        try:
            text = _extract(raw, ext)
        except Exception as exc:
            raise UnreadableDocument(
                f"not a readable {ext.lstrip('.')} file: {type(exc).__name__}: {exc}"
            ) from exc
    else:
        text = _extract(raw, ext)
    if len(text) > MAX_DOCUMENT_CHARS:
        return text[:MAX_DOCUMENT_CHARS] + _TRUNCATION_MARKER
    return text


def _read_docx(raw: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(raw))
    return "\n\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())


def _read_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    # pypdf narrates its own recovery at WARNING -- "invalid pdf header",
    # "EOF marker not found" -- and under the last-resort handler a command is
    # run with, each line lands in the terminal above the one that says what
    # varyz did about it. Whether the file could be read is reported here
    # either way: by a body of text, or by UnreadableDocument naming the file.
    noisy = logging.getLogger("pypdf")
    previous = noisy.level
    noisy.setLevel(logging.ERROR)
    try:
        reader = PdfReader(BytesIO(raw))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    finally:
        noisy.setLevel(previous)


def _extract(raw: bytes, ext: str) -> str:
    if ext == ".pdf":
        return _read_pdf(raw)
    if ext == ".docx":
        return _read_docx(raw)
    if ext == ".pptx":
        from pptx import Presentation

        deck = Presentation(BytesIO(raw))
        slides = []
        for index, slide in enumerate(deck.slides, 1):
            text = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if text:
                slides.append(f"Slide {index}\n" + "\n".join(text))
        return "\n\n".join(slides)
    if ext == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        blocks: list[str] = []
        try:
            for sheet in workbook.worksheets:
                output = io.StringIO()
                writer = csv.writer(output)
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index >= 5000:
                        writer.writerow(["… truncated …"])
                        break
                    writer.writerow(["" if value is None else value for value in row])
                blocks.append(f"Sheet: {sheet.title}\n{output.getvalue().strip()}")
        finally:
            workbook.close()
        return "\n\n".join(blocks)
    text = decode_text(raw)
    if ext in {".html", ".htm"}:
        return strip_html(text)
    if ext == ".json":
        try:
            rendered = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            # JSON accepts escaped lone surrogates, but SQLite and UTF-8 do not.
            # Keep their original escaped representation without losing data.
            return re.sub(r"[\ud800-\udfff]", lambda match: f"\\u{ord(match[0]):04x}", rendered)
        except json.JSONDecodeError:
            return text
    if ext == ".rtf":
        # Conservative plain-text extraction without executing embedded objects.
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", text)
        text = text.replace("{", " ").replace("}", " ")
        return re.sub(r"\s+", " ", text).strip()
    return text
